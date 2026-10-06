from pathlib import Path
from .base import BaseParser, ParsedDocument
from .ocr_utils import ocr_image
from config import OCR_ENABLED
from logger import get_logger

log = get_logger("parsers.office")


def _ocr_blob(blob: bytes, file_name: str = "") -> str:
    """이미지 blob → OCR 텍스트. 100×100px 미만 이미지(아이콘/로고)는 스킵."""
    try:
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        if img.width < 100 or img.height < 100:
            return ""
        return ocr_image(img)
    except Exception as e:
        log.warning(f"OCR 실패 ({file_name}): {e}")
        return ""


class DocxParser(BaseParser):
    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".docx"

    def parse(self, file_path: Path) -> ParsedDocument:
        from docx import Document
        try:
            doc = Document(str(file_path))
        except Exception:
            return ParsedDocument(content="", metadata={}, file_path=str(file_path))

        # 단락 텍스트
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        # 표 텍스트
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)

        # 이미지 OCR
        if OCR_ENABLED:
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        ocr_text = _ocr_blob(rel.target_part.blob, file_path.name)
                        if ocr_text.strip():
                            parts.append(ocr_text)
                    except Exception:
                        pass

        props = doc.core_properties
        metadata = {
            "title":  props.title  or "",
            "author": props.author or "",
        }
        return ParsedDocument(content="\n".join(parts), metadata=metadata, file_path=str(file_path))


class PptxParser(BaseParser):
    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".pptx"

    def parse(self, file_path: Path) -> ParsedDocument:
        from pptx import Presentation
        try:
            prs = Presentation(str(file_path))
        except Exception:
            return ParsedDocument(content="", metadata={}, file_path=str(file_path))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                # 텍스트 박스 / 제목 / 본문
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text)
                # 표 텍스트
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text_frame.text.strip()
                            for cell in row.cells
                            if cell.text_frame.text.strip()
                        )
                        if row_text:
                            parts.append(row_text)
                # 이미지 OCR
                # hasattr(shape, "image") 대신 try/except:
                # python-pptx는 embedded image 없을 때 ValueError를 raise하는데
                # hasattr은 AttributeError만 잡아 ValueError를 전파시킴
                if OCR_ENABLED:
                    try:
                        ocr_text = _ocr_blob(shape.image.blob, file_path.name)
                        if ocr_text.strip():
                            parts.append(ocr_text)
                    except Exception:
                        pass
        props = prs.core_properties
        metadata = {
            "title":       props.title  or "",
            "author":      props.author or "",
            "slide_count": len(prs.slides),
        }
        return ParsedDocument(content="\n".join(parts), metadata=metadata, file_path=str(file_path))


class XlsxParser(BaseParser):
    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in (".xlsx", ".xls")

    def parse(self, file_path: Path) -> ParsedDocument:
        import math
        import pandas as pd
        from .base import ParsedTable

        engine = "xlrd" if file_path.suffix.lower() == ".xls" else "openpyxl"
        try:
            sheets = pd.read_excel(str(file_path), sheet_name=None, engine=engine)
        except Exception:
            return ParsedDocument(content="", metadata={}, file_path=str(file_path))
        if not sheets:
            return ParsedDocument(content="", metadata={}, file_path=str(file_path))

        def _clean(v):
            if v is None:
                return None
            try:
                if isinstance(v, float) and math.isnan(v):
                    return None
            except (TypeError, ValueError):
                pass
            return v

        tables = []
        for sheet_name, df in sheets.items():
            try:
                df = df.dropna(how="all").reset_index(drop=True)
                headers = [str(col) for col in df.columns]
                rows = [[_clean(v) for v in row] for row in df.values.tolist()]
                tables.append(ParsedTable(
                    sheet_name=str(sheet_name),
                    headers=headers,
                    rows=rows,
                ))
            except Exception:
                pass

        metadata = {
            "sheet_count": len(sheets),
            "sheet_names": ", ".join(str(k) for k in sheets.keys()),
        }
        return ParsedDocument(tables=tables, metadata=metadata, file_path=str(file_path))
